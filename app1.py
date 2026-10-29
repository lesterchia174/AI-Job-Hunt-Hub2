"""
AI Job Hunt OS — Test Build v4.2.3.3 (Supabase Edition)
Streamlit app that tracks job applications and uses AI (Groq + AI Singapore)
to tailor resumes and generate cover letters.

Now with:
- Quick Import from Job Link inside the existing "New Application" view
- MyCareersFuture Job Search (search → preview → one-click pre-fill of New Application)
- "Edit & Export" view with a rich text (Markdown) editor (streamlit-lexical-extended)
- PDF / DOCX export, downloaded to the user's local drive (no .txt, no Storage bucket)
- Filenames written back to the existing `applications` table columns
  (resume_filename / cover_letter_filename / doc_number)
- LLM Monitoring (tokens, cost, latency, failures) including cached input tokens
- Human-in-the-Loop document generation choices
- Document naming format: PWA00286_Resume for [Role]_[Company]_DDMMYY_CV

Run:
    streamlit run app.py
"""

import hashlib
import io
import os
import re
import time
from datetime import datetime
from xml.sax.saxutils import escape as xml_escape
from zoneinfo import ZoneInfo

import pandas as pd
import requests
from bs4 import BeautifulSoup
import streamlit as st
from dotenv import load_dotenv
from sqlalchemy import Column, DateTime, Float, Integer, String, Text, create_engine
from sqlalchemy.orm import declarative_base, sessionmaker
from groq import Groq
from openai import OpenAI

# MyCareersFuture search adapter (optional — app degrades gracefully if absent)
try:
    from job_search import search_mycareersfuture, JobPosting
    _JOB_SEARCH_AVAILABLE = True
except ImportError:
    search_mycareersfuture = None
    JobPosting = None
    _JOB_SEARCH_AVAILABLE = False

# Optional UI components (the app degrades gracefully if they are missing)
try:
    from streamlit_lexical_extended import streamlit_lexical_extended
except ImportError:
    streamlit_lexical_extended = None

# ---------------------------------------------------------------------------
# 1. CONFIGURATION
# ---------------------------------------------------------------------------

st.set_page_config(
    page_title="AI Job Hunt OS",
    page_icon="💼",
    layout="wide",
    initial_sidebar_state="expanded",
)

load_dotenv()

# ============================================================
# 🚀 API KEY & DATABASE CONFIGURATION
# ============================================================

def _get_secret(key: str, default=None):
    try:
        from google.colab import userdata
        val = userdata.get(key)
        if val:
            return val
    except ImportError:
        pass

    val = os.environ.get(key)
    if val:
        return val

    try:
        val = st.secrets.get(key)
        if val:
            return val
    except Exception:
        pass

    return default


GROQ_API_KEY = _get_secret("GROQ_API_KEY")
SEALION_API_KEY = _get_secret("SEALION_API_KEY")
HF_TOKEN = _get_secret("HF_TOKEN")
SUPABASE_DB_URL = _get_secret("SUPABASE_DB_URL")

if not GROQ_API_KEY and not SEALION_API_KEY:
    st.error(
        "❌ No API keys found. Set at least one of GROQ_API_KEY or "
        "SEALION_API_KEY in your `.env` file or Streamlit secrets."
    )
    st.stop()

if not SUPABASE_DB_URL:
    st.error(
        "❌ No SUPABASE_DB_URL found. Set your Supabase PostgreSQL connection string "
        "in your `.env` file or Streamlit secrets."
    )
    st.stop()

if HF_TOKEN:
    os.environ["HF_TOKEN"] = HF_TOKEN

groq_client = Groq(api_key=GROQ_API_KEY) if GROQ_API_KEY else None

AI_SINGAPORE_BASE_URL = "https://api.sea-lion.ai/v1"
ai_singapore_client = (
    OpenAI(api_key=SEALION_API_KEY, base_url=AI_SINGAPORE_BASE_URL)
    if SEALION_API_KEY else None
)

# Constants -----------------------------------------------------------------

STAGES = [
    {"id": "Applied",      "name": "Applied",      "color": "#3498db"},
    {"id": "Acknowledged", "name": "Acknowledged", "color": "#f39c12"},
    {"id": "Screening",    "name": "Screening",    "color": "#9b59b6"},
    {"id": "Interviewing", "name": "Interviewing", "color": "#e67e22"},
    {"id": "Offer",        "name": "Offer",        "color": "#27ae60"},
    {"id": "Rejected",     "name": "Rejected",     "color": "#e74c3c"},
]

# ============================================================
# 🔌 DATABASE URL & DRIVER SETUP (OPTION 1: psycopg v3)
# ============================================================

if SUPABASE_DB_URL.startswith("postgres://"):
    DATABASE_URL = SUPABASE_DB_URL.replace("postgres://", "postgresql+psycopg://", 1)
elif SUPABASE_DB_URL.startswith("postgresql://"):
    DATABASE_URL = SUPABASE_DB_URL.replace("postgresql://", "postgresql+psycopg://", 1)
else:
    DATABASE_URL = SUPABASE_DB_URL

# ============================================================
# 🧠 MODEL REGISTRY
# ============================================================

MODEL_REGISTRY = {
    "Sea-Lion v4 27B": {
        "provider": "ai_singapore",
        "model_id": "aisingapore/Gemma-SEA-LION-v4-27B-IT",
        "description": "Gemma 3 27B fine-tuned for SEA languages by AI Singapore.",
    },
    "Sea-Lion v4 32B (Qwen)": {
        "provider": "ai_singapore",
        "model_id": "aisingapore/Qwen-SEA-LION-v4-32B-IT",
        "description": "Qwen3-32B fine-tuned for SEA tasks.",
    },
    "Groq: GPT-OSS 120B": {
        "provider": "groq",
        "model_id": "openai/gpt-oss-120b",
        "description": "OpenAI's open-weight MoE, 120B params, reasoning-capable.",
    },
    "Groq: GPT-OSS 20B": {
        "provider": "groq",
        "model_id": "openai/gpt-oss-20b",
        "description": "Smaller, faster OpenAI open-weight model.",
    },
}

DEFAULT_MODEL_KEY = "Sea-Lion v4 27B"

# ============================================================
# 💰 MODEL PRICING (per 1M tokens, USD)
# ============================================================

MODEL_PRICING = {
    "Qwen/Qwen3.5-9B": {
        "input": 0.0, "cached_input": 0.0, "output": 0.0,
        "context_window": 131072, "max_output": 32768,
    },
    "openai/gpt-oss-120b": {
        "input": 0.15, "cached_input": 0.075, "output": 0.60,
        "context_window": 131072, "max_output": 65536,
    },
    "openai/gpt-oss-20b": {
        "input": 0.075, "cached_input": 0.037, "output": 0.30,
        "context_window": 131072, "max_output": 65536,
    },
    "aisingapore/Gemma-SEA-LION-v4-27B-IT": {
        "input": 0.0, "cached_input": 0.0, "output": 0.0,
        "context_window": 131072, "max_output": 65536,
    },
    "aisingapore/Qwen-SEA-LION-v4-32B-IT": {
        "input": 0.0, "cached_input": 0.0, "output": 0.0,
        "context_window": 131072, "max_output": 65536,
    },
}


def get_pricing(model_id: str) -> dict:
    return MODEL_PRICING.get(
        model_id,
        {"input": 0.0, "output": 0.0, "cached_input": 0.0,
         "context_window": 0, "max_output": 0},
    )


def compute_cost_usd(
    model_id: str,
    uncached_input_tokens: int,
    output_tokens: int,
    cached_input_tokens: int = 0,
) -> float:
    p = get_pricing(model_id)
    return (
        (uncached_input_tokens / 1_000_000) * p["input"]
        + (cached_input_tokens / 1_000_000) * p["cached_input"]
        + (output_tokens / 1_000_000) * p["output"]
    )


# ---------------------------------------------------------------------------
# 2. DATABASE LAYER
# ---------------------------------------------------------------------------

Base = declarative_base()


class Application(Base):
    __tablename__ = "applications"

    id = Column(Integer, primary_key=True)
    doc_number = Column(Integer, default=0)
    company = Column(String, nullable=False)
    role = Column(String, nullable=False)
    job_url = Column(String, default="")
    status = Column(String, default="Applied")
    job_description = Column(Text, default="")
    resume_text = Column(Text, default="")
    cover_letter = Column(Text, default="")
    resume_filename = Column(String, default="")
    cover_letter_filename = Column(String, default="")
    match_score = Column(String, default="")
    applied_date = Column(DateTime, default=datetime.now)
    last_updated = Column(DateTime, default=datetime.now, onupdate=datetime.now)

    def to_dict(self):
        return {c.name: getattr(self, c.name) for c in self.__table__.columns}


class LLMUsage(Base):
    __tablename__ = "llm_usage"

    id = Column(Integer, primary_key=True)
    called_at = Column(DateTime, default=datetime.now, index=True)
    provider = Column(String)
    model_key = Column(String)
    model_id = Column(String)
    operation = Column(String)
    input_tokens = Column(Integer, default=0)
    cached_input_tokens = Column(Integer, default=0)
    output_tokens = Column(Integer, default=0)
    total_tokens = Column(Integer, default=0)
    latency_seconds = Column(Float, default=0.0)
    cost_usd = Column(Float, default=0.0)
    success = Column(Integer, default=1)
    error_message = Column(Text, default="")

    def to_dict(self):
        return {c.name: getattr(self, c.name) for c in self.__table__.columns}


@st.cache_resource
def get_engine():
    engine = create_engine(
        DATABASE_URL,
        pool_pre_ping=True,
        pool_recycle=120,
        pool_size=5,
        max_overflow=5,
    )
    return engine


@st.cache_resource
def get_session_factory():
    return sessionmaker(bind=get_engine(), autoflush=False, autocommit=False)


def db_session():
    return get_session_factory()()


# ---------------------------------------------------------------------------
# 3. AI & SCRAPING LAYER
# ---------------------------------------------------------------------------

@st.cache_resource(show_spinner="Loading embedding model (first run only)...")
def load_embedding_model():
    try:
        from sentence_transformers import SentenceTransformer
        return SentenceTransformer("all-MiniLM-L6-v2")
    except Exception as e:
        st.warning(f"Embedding model unavailable: {e}")
        return None


def fetch_job_details_from_url(url: str) -> dict:
    """
    Scrapes a job URL to extract company name, job title, and job description.
    Returns a dict with: company, role, job_description, job_url.
    """
    result = {"company": "", "role": "", "job_description": "", "job_url": url}

    if not url or not url.startswith("http"):
        return result

    try:
        headers = {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/120.0.0.0 Safari/537.36"
            )
        }
        response = requests.get(url, headers=headers, timeout=10)
        if response.status_code != 200:
            return result

        soup = BeautifulSoup(response.text, "html.parser")

        og_title = soup.find("meta", property="og:title")
        if og_title and og_title.get("content"):
            result["role"] = og_title["content"].strip()

        og_site = soup.find("meta", property="og:site_name")
        if og_site and og_site.get("content"):
            result["company"] = og_site["content"].strip()

        if not result["role"] and soup.title and soup.title.string:
            result["role"] = soup.title.string.strip()

        for script in soup(["script", "style", "nav", "footer", "header", "noscript", "svg"]):
            script.extract()

        text = soup.get_text(separator="\n", strip=True)
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        result["job_description"] = "\n".join(lines[:150])

    except Exception as e:
        print(f"Error fetching URL: {e}")

    return result


def get_active_model() -> tuple[str, str, str]:
    key = st.session_state.get("selected_model_key", DEFAULT_MODEL_KEY)
    entry = MODEL_REGISTRY.get(key, MODEL_REGISTRY[DEFAULT_MODEL_KEY])
    return entry["provider"], entry["model_id"], key


def load_llm_for(provider: str, model_id: str):
    if provider == "groq":
        if not GROQ_API_KEY:
            return None, "Groq API key not set"
        try:
            from langchain_groq import ChatGroq
            llm = ChatGroq(model=model_id, temperature=0.7, api_key=GROQ_API_KEY)
            return llm, f"Groq ({model_id})"
        except Exception as e:
            return None, f"Groq init failed: {e}"

    if provider == "ai_singapore":
        api_key_to_use = SEALION_API_KEY or HF_TOKEN
        base_url_to_use = (
            AI_SINGAPORE_BASE_URL if SEALION_API_KEY
            else "https://api-inference.huggingface.co/v1/"
        )
        if not api_key_to_use:
            return None, "API key or HF_TOKEN not set for Hugging Face / Sea-Lion endpoint"
        try:
            from langchain_openai import ChatOpenAI
            llm = ChatOpenAI(
                model=model_id,
                temperature=0.7,
                api_key=api_key_to_use,
                base_url=base_url_to_use,
            )
            return llm, f"Provider ({model_id})"
        except Exception as e:
            return None, f"Provider init failed: {e}"

    return None, f"Unknown provider: {provider}"


def _extract_usage_from_langchain(result) -> tuple[int, int, int]:
    um = getattr(result, "usage_metadata", None)
    if um:
        input_tokens = int(um.get("input_tokens", 0) or 0)
        output_tokens = int(um.get("output_tokens", 0) or 0)
        details = um.get("input_token_details") or {}
        cached = int(
            details.get("cache_read", 0)
            or details.get("cached_tokens", 0)
            or 0
        )
        return input_tokens, cached, output_tokens

    rm = getattr(result, "response_metadata", None) or {}
    tu = rm.get("token_usage") or rm.get("usage") or {}
    input_tokens = int(tu.get("prompt_tokens", 0) or 0)
    output_tokens = int(tu.get("completion_tokens", 0) or 0)
    details = tu.get("prompt_tokens_details") or {}
    cached = int(details.get("cached_tokens", 0) or 0)
    return input_tokens, cached, output_tokens


def _extract_cached_tokens_from_usage(usage) -> int:
    if not usage:
        return 0
    details = getattr(usage, "prompt_tokens_details", None)
    if details is None and isinstance(usage, dict):
        details = usage.get("prompt_tokens_details") or {}
    if isinstance(details, dict):
        return int(details.get("cached_tokens", 0) or 0)
    return int(getattr(details, "cached_tokens", 0) or 0)


def svc_log_llm_usage(
    *, provider, model_key, model_id, operation,
    input_tokens, output_tokens, latency_seconds,
    cached_input_tokens=0, success=True, error_message="",
):
    total_prompt = int(input_tokens or 0)
    cached = int(cached_input_tokens or 0)
    uncached = max(total_prompt - cached, 0)
    output = int(output_tokens or 0)

    cost = (
        compute_cost_usd(model_id, uncached, output, cached)
        if success else 0.0
    )

    with db_session() as db:
        row = LLMUsage(
            provider=provider,
            model_key=model_key,
            model_id=model_id,
            operation=operation,
            input_tokens=uncached,
            cached_input_tokens=cached,
            output_tokens=output,
            total_tokens=uncached + cached + output,
            latency_seconds=float(latency_seconds or 0.0),
            cost_usd=cost,
            success=1 if success else 0,
            error_message=error_message or "",
        )
        db.add(row)
        db.commit()


def calculate_match_score(resume_text: str, job_description: str) -> float:
    if not resume_text or not job_description:
        return 0.0
    model = load_embedding_model()
    if model is None:
        return 0.0
    from sentence_transformers import util
    emb = model.encode([job_description, resume_text], convert_to_tensor=True)
    score = util.cos_sim(emb[0], emb[1])
    return float(score[0][0])


def call_llm(prompt: str, operation: str = "generic") -> str:
    provider, model_id, label = get_active_model()
    llm, status = load_llm_for(provider, model_id)

    if llm is None:
        svc_log_llm_usage(
            provider=provider,
            model_key=label,
            model_id=model_id,
            operation=operation,
            input_tokens=0,
            cached_input_tokens=0,
            output_tokens=0,
            latency_seconds=0.0,
            success=False,
            error_message=str(status),
        )
        return f"⚠️ {status}"

    start = time.time()
    try:
        result = llm.invoke(prompt)
        latency = time.time() - start
        in_tok, cached_tok, out_tok = _extract_usage_from_langchain(result)

        svc_log_llm_usage(
            provider=provider,
            model_key=label,
            model_id=model_id,
            operation=operation,
            input_tokens=in_tok,
            cached_input_tokens=cached_tok,
            output_tokens=out_tok,
            latency_seconds=latency,
            success=True,
        )
        return getattr(result, "content", str(result))
    except Exception as e:
        latency = time.time() - start
        svc_log_llm_usage(
            provider=provider,
            model_key=label,
            model_id=model_id,
            operation=operation,
            input_tokens=0,
            cached_input_tokens=0,
            output_tokens=0,
            latency_seconds=latency,
            success=False,
            error_message=str(e),
        )
        return f"⚠️ LLM call failed ({status}): {e}"


def call_groq_direct(prompt: str, operation: str = "direct_test") -> str:
    if groq_client is None:
        return "⚠️ Groq client not initialized (no API key)."
    provider, model_id, label = get_active_model()
    if provider != "groq":
        return "⚠️ Currently selected model is not a Groq model."

    start = time.time()
    try:
        resp = groq_client.chat.completions.create(
            model=model_id,
            messages=[{"role": "user", "content": prompt}],
            temperature=0.7,
        )
        latency = time.time() - start

        usage = getattr(resp, "usage", None)
        in_tok = int(getattr(usage, "prompt_tokens", 0) or 0)
        out_tok = int(getattr(usage, "completion_tokens", 0) or 0)
        cached_tok = _extract_cached_tokens_from_usage(usage)

        svc_log_llm_usage(
            provider="groq",
            model_key=label,
            model_id=model_id,
            operation=operation,
            input_tokens=in_tok,
            cached_input_tokens=cached_tok,
            output_tokens=out_tok,
            latency_seconds=latency,
            success=True,
        )
        return resp.choices[0].message.content
    except Exception as e:
        latency = time.time() - start
        svc_log_llm_usage(
            provider="groq",
            model_key=label,
            model_id=model_id,
            operation=operation,
            input_tokens=0,
            cached_input_tokens=0,
            output_tokens=0,
            latency_seconds=latency,
            success=False,
            error_message=str(e),
        )
        return f"⚠️ Groq direct call failed ({model_id}): {e}"


def call_ai_singapore_direct(prompt: str, operation: str = "direct_test") -> str:
    if ai_singapore_client is None:
        return "⚠️ AI Singapore client not initialized (no API key)."
    provider, model_id, label = get_active_model()
    if provider != "ai_singapore":
        return "⚠️ Currently selected model is not an AI Singapore model."

    start = time.time()
    try:
        resp = ai_singapore_client.chat.completions.create(
            model=model_id,
            messages=[{"role": "user", "content": prompt}],
            temperature=0.7,
        )
        latency = time.time() - start

        usage = getattr(resp, "usage", None)
        in_tok = int(getattr(usage, "prompt_tokens", 0) or 0)
        out_tok = int(getattr(usage, "completion_tokens", 0) or 0)
        cached_tok = _extract_cached_tokens_from_usage(usage)

        svc_log_llm_usage(
            provider="ai_singapore",
            model_key=label,
            model_id=model_id,
            operation=operation,
            input_tokens=in_tok,
            cached_input_tokens=cached_tok,
            output_tokens=out_tok,
            latency_seconds=latency,
            success=True,
        )
        return resp.choices[0].message.content
    except Exception as e:
        latency = time.time() - start
        svc_log_llm_usage(
            provider="ai_singapore",
            model_key=label,
            model_id=model_id,
            operation=operation,
            input_tokens=0,
            cached_input_tokens=0,
            output_tokens=0,
            latency_seconds=latency,
            success=False,
            error_message=str(e),
        )
        return f"⚠️ AI Singapore direct call failed ({model_id}): {e}"


def tailor_resume(resume_text: str, job_description: str) -> str:
    prompt = f"""You are an objective professional resume writer. Your task is to draft a resume based ONLY on the provided job description.
Strict Rules:
Use quantifiable achievements and keywords from the job posting.
Do not invent companies, job titles, dates, or metrics (e.g., percentages, dollar amounts) that are missing from the raw text.
If a metric is not explicitly stated in the source text, describe the qualitative impact instead of guessing a number.
If information is missing, use placeholders like [Insert Metric] rather than fabricating data.
Keep the format clean and professional. Do NOT invent experience.


JOB DESCRIPTION:
{job_description}

CURRENT RESUME:
{resume_text}

TAILORED RESUME:"""
    return call_llm(prompt, operation="tailor_resume")


def generate_cover_letter(resume_text: str, job_description: str, company: str, role: str) -> str:
    prompt = f"""Write a professional, concise cover letter for the {role} role at {company}.
Ground it in the resume and job description below. Do NOT use placeholders like [Your Name].

JOB DESCRIPTION:
{job_description}

RESUME:
{resume_text}

COVER LETTER:"""
    return call_llm(prompt, operation="cover_letter")


def extract_pdf_text(uploaded_file) -> str:
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(uploaded_file)
        return "".join((page.extract_text() or "") for page in reader.pages)
    except Exception as e:
        st.error(f"Failed to read PDF: {e}")
        return ""


def classify_email_status(body: str) -> str:
    b = body.lower()
    if any(k in b for k in ["unfortunately", "not moving forward", "other candidates", "rejected"]):
        return "Rejected"
    if any(k in b for k in ["interview", "schedule a call", "speak with"]):
        return "Interviewing"
    if any(k in b for k in ["offer", "pleased to offer"]):
        return "Offer"
    if any(k in b for k in ["received your application", "thank you for applying", "acknowledge"]):
        return "Acknowledged"
    return "Applied"


# ---------------------------------------------------------------------------
# 3b. DOCUMENT NAMING HELPERS
# ---------------------------------------------------------------------------

DOC_NUMBER_START = 291


def sanitize_filename_part(value: str, max_len: int = 60) -> str:
    value = (value or "").strip()
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "", value)
    value = re.sub(r"\s+", " ", value)
    return value.strip(" .")[:max_len].strip(" .")


def clean_text(s) -> str:
    return re.sub(r"\*+", "", str(s or "")).strip()


def _today_ddmmyy() -> str:
    try:
        now = datetime.now(ZoneInfo("Asia/Singapore"))
    except Exception:
        now = datetime.now()
    return now.strftime("%d%m%y")


def get_next_doc_number() -> int:
    from sqlalchemy import func
    with db_session() as db:
        max_doc = db.query(func.max(Application.doc_number)).scalar()
    if max_doc:
        return max(int(max_doc) + 1, DOC_NUMBER_START)
    return DOC_NUMBER_START


def build_document_filename(
    doc_number: int,
    role: str,
    company: str,
    doc_type: str,
) -> str:
    date_str = _today_ddmmyy()
    role_part = sanitize_filename_part(role)
    company_part = sanitize_filename_part(company)
    prefix = f"PWA{int(doc_number):05d}"

    if doc_type.upper() == "CV":
        return f"{prefix}_Resume for {role_part}_{company_part}_{date_str}_CV"

    return f"{prefix}_Cover Letter for {role_part}_{company_part}_{date_str}_CL"


# ---------------------------------------------------------------------------
# 3c. EXPORT HELPERS (Markdown -> PDF / DOCX)
# ---------------------------------------------------------------------------

MIME_PDF = "application/pdf"
MIME_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _unescape_md(text: str) -> str:
    return re.sub(r"\\([\\`*_{}\[\]()#+\-.!|>~])", r"\1", text)


def _parse_markdown_blocks(md: str):
    blocks = []
    in_skills_or_competencies = False
    lines = (md or "").replace("\r\n", "\n").split("\n")

    i = 0
    while i < len(lines):
        raw = lines[i]
        line = re.sub(r"(<br\s*/?>|\\)\s*$", "", raw.rstrip()).rstrip()

        if not line.strip():
            in_skills_or_competencies = False
            blocks.append(("blank", ""))
            i += 1
            continue

        # --- TABLE DETECTION ---
        if line.startswith("|") and line.endswith("|"):
            if i + 1 < len(lines):
                next_line = lines[i + 1].strip()
                if re.match(r"^\|[\s\-:|]+\|$", next_line):
                    headers = [c.strip() for c in line.strip("|").split("|")]
                    rows = []
                    i += 2
                    while i < len(lines):
                        row_line = lines[i].strip()
                        if not (row_line.startswith("|") and row_line.endswith("|")):
                            break
                        row_cells = [c.strip() for c in row_line.strip("|").split("|")]
                        rows.append(row_cells)
                        i += 1
                    blocks.append(("table", {"headers": headers, "rows": rows}))
                    continue
        # -----------------------

        if re.match(r"^\s*([-*_])\1{2,}\s*$", line):
            blocks.append(("blank", ""))
            i += 1
            continue

        m_head = re.match(r"^(#{1,3})\s+(.*)$", line)
        if m_head:
            heading_text = m_head.group(2).replace("**", "").strip().upper()
            if ("CORE COMPETENCIES" in heading_text
                    or "TECHNICAL SKILLSETS" in heading_text
                    or "EDUCATION" in heading_text):
                in_skills_or_competencies = True
            else:
                in_skills_or_competencies = False
            blocks.append((
                f"h{len(m_head.group(1))}",
                _unescape_md(m_head.group(2).replace("**", "").strip()),
            ))
            i += 1
            continue

        m_bullet = re.match(r"^\s*[-*•]\s+(.*)$", line)
        if m_bullet:
            blocks.append(("bullet", m_bullet.group(1).strip()))
            i += 1
            continue

        if in_skills_or_competencies:
            is_job_header = (
                bool(re.search(
                    r"(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec|\b20\d{2}\b)",
                    line,
                ))
                or line.endswith(", Singapore")
            )
            if not is_job_header:
                blocks.append(("bullet", line.strip()))
                i += 1
                continue

        blocks.append(("para", line.strip()))
        i += 1

    return blocks


def _inline_tokens(text: str):
    for part in re.split(r"(\*\*.+?\*\*|\*[^*\s][^*]*?\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            yield _unescape_md(part[2:-2]), True, False
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            yield _unescape_md(part[1:-1]), False, True
        else:
            yield _unescape_md(part), False, False


def markdown_to_docx_bytes(md: str) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    for item in _parse_markdown_blocks(md):
        kind = item[0]
        if kind == "blank":
            continue

        if kind == "table":
            data = item[1]
            headers = data["headers"]
            rows = data["rows"]

            table = doc.add_table(rows=1, cols=len(headers))
            table.style = "Table Grid"

            hdr_cells = table.rows[0].cells
            for j, header in enumerate(headers):
                para = hdr_cells[j].paragraphs[0]
                clean_header = re.sub(r"<br\s*/?>", " ", header)
                for t, bold, italic in _inline_tokens(clean_header):
                    run = para.add_run(t)
                    run.bold = True
                    run.italic = italic

            for row_data in rows:
                row_cells = table.add_row().cells
                for j, cell_data in enumerate(row_data):
                    if j >= len(headers):
                        break
                    para = row_cells[j].paragraphs[0]
                    clean_cell = re.sub(r"<br\s*/?>", " ", cell_data)
                    for t, bold, italic in _inline_tokens(clean_cell):
                        run = para.add_run(t)
                        run.bold = bold
                        run.italic = italic
            continue

        text = item[1]
        if kind in ("h1", "h2", "h3"):
            doc.add_heading(text, level=int(kind[1]))
            continue
        para = (
            doc.add_paragraph(style="List Bullet")
            if kind == "bullet" else doc.add_paragraph()
        )
        for t, bold, italic in _inline_tokens(text):
            run = para.add_run(t)
            run.bold = bold
            run.italic = italic

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def clean_pdf_text(text: str) -> str:
    if not text:
        return ""
    text = text.replace("\u2011", "-").replace("\u2013", "-").replace("\u2014", "-")
    text = text.replace("\u00a0", " ")
    return text


def markdown_to_pdf_bytes(md: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from reportlab.lib import colors

    # --- RIGOROUS UNICODE SANITIZATION ---
    if md:
        md = (
            md.replace("\u00a0", " ")
              .replace("\u202f", " ")
              .replace("\u2002", " ")
              .replace("\u2003", " ")
              .replace("\u2011", "-")
              .replace("\u2013", "-")
              .replace("\u2014", "-")
              .replace("\u201c", '"')
              .replace("\u201d", '"')
              .replace("\u2018", "'")
              .replace("\u2019", "'")
        )
    # ------------------------------------

    buf = io.BytesIO()
    pdf = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=2 * cm, rightMargin=2 * cm,
        topMargin=2 * cm, bottomMargin=2 * cm,
    )
    base = getSampleStyleSheet()
    body = ParagraphStyle("Body", parent=base["BodyText"],
                          fontName="Helvetica", fontSize=10.5, leading=14)
    heads = {
        "h1": ParagraphStyle("H1", parent=body, fontName="Helvetica-Bold",
                             fontSize=16, leading=20, spaceAfter=6),
        "h2": ParagraphStyle("H2", parent=body, fontName="Helvetica-Bold",
                             fontSize=13, leading=17, spaceBefore=6, spaceAfter=4),
        "h3": ParagraphStyle("H3", parent=body, fontName="Helvetica-Bold",
                             fontSize=11.5, leading=15, spaceBefore=4, spaceAfter=2),
    }
    bullet_style = ParagraphStyle(
        "Bullet", parent=body,
        leftIndent=18, bulletIndent=6, spaceAfter=3,
    )
    table_cell_style = ParagraphStyle(
        "TableCell", parent=body, fontSize=9, leading=12,
    )

    def to_markup(text: str) -> str:
        cleaned_text = re.sub(r"^\s*[-*•]\s+", "", text)
        cleaned_text = re.sub(r"<br\s*/?>", " ", cleaned_text)
        out = []
        for t, bold, italic in _inline_tokens(cleaned_text):
            s = xml_escape(t)
            if bold:
                s = f"<b>{s}</b>"
            if italic:
                s = f"<i>{s}</i>"
            out.append(s)
        return "".join(out)

    story = []
    for item in _parse_markdown_blocks(md):
        kind = item[0]
        if kind == "blank":
            story.append(Spacer(1, 6))
        elif kind == "table":
            data = item[1]
            headers = data["headers"]
            rows = data["rows"]

            table_data = []
            header_row = [Paragraph(to_markup(h), table_cell_style) for h in headers]
            table_data.append(header_row)

            for row in rows:
                row_cells = []
                for i in range(len(headers)):
                    cell_text = row[i] if i < len(row) else ""
                    row_cells.append(Paragraph(to_markup(cell_text), table_cell_style))
                table_data.append(row_cells)

            avail_width = 17 * cm
            col_width = avail_width / len(headers) if headers else avail_width
            col_widths = [col_width] * len(headers)

            t = Table(table_data, colWidths=col_widths)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.black),
                ("BOX", (0, 0), (-1, -1), 0.25, colors.black),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(t)
            story.append(Spacer(1, 6))
        elif kind in heads:
            story.append(Paragraph(xml_escape(item[1]), heads[kind]))
        elif kind == "bullet":
            story.append(Paragraph(to_markup(item[1]), bullet_style, bulletText="•"))
        else:
            story.append(Paragraph(to_markup(item[1]), body))
    if not story:
        story.append(Paragraph(" ", body))

    pdf.build(story)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# 4. SERVICE LAYER
# ---------------------------------------------------------------------------

def svc_list_applications():
    with db_session() as db:
        return [
            a.to_dict()
            for a in db.query(Application).order_by(Application.last_updated.desc()).all()
        ]


def svc_get_application(app_id: int):
    with db_session() as db:
        app = db.query(Application).filter(Application.id == app_id).first()
        return app.to_dict() if app else None


def svc_create_application(**kwargs):
    with db_session() as db:
        app = Application(**kwargs)
        db.add(app)
        db.commit()
        db.refresh(app)
        return app.id


def svc_update_application(app_id: int, **fields) -> bool:
    with db_session() as db:
        app = db.query(Application).filter(Application.id == app_id).first()
        if not app:
            return False
        for k, v in fields.items():
            setattr(app, k, v)
        app.last_updated = datetime.now()
        db.commit()
        return True


def svc_update_status(app_id: int, new_status: str):
    with db_session() as db:
        app = db.query(Application).filter(Application.id == app_id).first()
        if app:
            app.status = new_status
            app.last_updated = datetime.now()
            db.commit()
            return True
    return False


def svc_delete_application(app_id: int):
    with db_session() as db:
        app = db.query(Application).filter(Application.id == app_id).first()
        if app:
            db.delete(app)
            db.commit()


def svc_list_llm_usage(days: int = 7):
    from datetime import timedelta
    cutoff = datetime.now() - timedelta(days=days)
    with db_session() as db:
        rows = (
            db.query(LLMUsage)
            .filter(LLMUsage.called_at >= cutoff)
            .order_by(LLMUsage.called_at.desc())
            .all()
        )
        return [r.to_dict() for r in rows]


# ---------------------------------------------------------------------------
# 5. UI LAYER
# ---------------------------------------------------------------------------

def render_model_selector():
    st.sidebar.markdown("### 🧠 AI Model")
    options = list(MODEL_REGISTRY.keys())
    default_idx = options.index(DEFAULT_MODEL_KEY) if DEFAULT_MODEL_KEY in options else 0

    selected = st.sidebar.selectbox(
        "Select model", options,
        index=st.session_state.get("selected_model_index", default_idx),
        key="selected_model_widget",
        label_visibility="collapsed",
    )
    st.session_state["selected_model_key"] = selected
    st.session_state["selected_model_index"] = options.index(selected)

    entry = MODEL_REGISTRY[selected]
    provider = entry["provider"]
    st.sidebar.caption(f"Provider: **{provider}**")
    st.sidebar.caption(f"`{entry['model_id']}`")
    st.sidebar.caption(entry["description"])

    if provider == "groq" and not GROQ_API_KEY:
        st.sidebar.error("Groq key missing.")
    if provider == "ai_singapore" and not SEALION_API_KEY:
        st.sidebar.error("AI Singapore (Sea-Lion) key missing.")


def render_sidebar():
    with st.sidebar:
        st.title("💼 AI Job Hunt Hub")
        st.caption("AI-powered job search tracker")
        st.markdown("---")

        st.markdown("**👤 Lester**")
        st.caption("lesterchia1@hotmail.com")
        st.markdown("---")

        nav = st.radio(
            "Navigation",
            ["📊 Tracker",
             "🔎 Job Search",
             "➕ New Application",
             "✏️ Edit & Export",
             "📥 Outbox",
             "📈 Analytics",
             "🔍 Monitoring",
             "⚙️ Settings"],
            index=0,
            key="nav_radio",
            label_visibility="collapsed",
        )
        st.markdown("---")

        search = st.text_input("🔍 Search company / role", "")

        st.markdown("---")
        render_model_selector()
        st.markdown("---")

        provider, model_id, label = get_active_model()
        llm, status = load_llm_for(provider, model_id)
        if llm is None:
            st.error(f"LLM unavailable: {status}")
        else:
            st.success(f"LLM ready: {label}")

    return nav, search


def render_kpi_banner(df: pd.DataFrame):
    total = len(df)
    applied = int((df["status"] == "Applied").sum()) if total else 0
    interviewing = int((df["status"] == "Interviewing").sum()) if total else 0
    offers = int((df["status"] == "Offer").sum()) if total else 0

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Total Applications", total)
    c2.metric("Applied", applied)
    c3.metric("Interviewing", interviewing)
    c4.metric("Offers", offers)


def render_kanban(df: pd.DataFrame):
    stage_ids = [s["id"] for s in STAGES]
    cols = st.columns(len(STAGES))
    for col, stage in zip(cols, STAGES):
        stage_df = df[df["status"] == stage["id"]]
        with col:
            st.markdown(
                f"""<div style="background:{stage['color']};color:white;
                padding:6px 10px;border-radius:6px;font-weight:600;
                font-size:13px;text-align:center;">
                {stage['name']} ({len(stage_df)})</div>""",
                unsafe_allow_html=True,
            )
            st.write("")

            if stage_df.empty:
                st.caption("—")
                continue

            for _, row in stage_df.iterrows():
                with st.container(border=True):
                    st.markdown(f"**{clean_text(row['role'])}**")
                    st.caption(f"🏢 {clean_text(row['company'])}")
                    if row.get("match_score"):
                        st.caption(f"🎯 {row['match_score']}")

                    new_status = st.selectbox(
                        "Move to",
                        stage_ids,
                        index=stage_ids.index(row["status"]),
                        key=f"status_{row['id']}_{row['status']}",
                        label_visibility="collapsed",
                    )
                    if new_status != row["status"]:
                        svc_update_status(int(row["id"]), new_status)
                        st.rerun()

                    with st.expander("Details"):
                        st.write(f"**Applied:** {row['applied_date']}")
                        if row.get("job_url"):
                            st.markdown(f"[Job posting]({row['job_url']})")

                        if row.get("resume_filename"):
                            st.caption(f"📄 Resume: `{row['resume_filename']}`")
                        if row.get("cover_letter_filename"):
                            st.caption(f"✉️ Cover: `{row['cover_letter_filename']}`")

                        if st.button("🗑 Delete", key=f"del_{row['id']}"):
                            svc_delete_application(int(row["id"]))
                            st.rerun()


def view_tracker(search: str):
    st.title("📊 Tracker")
    st.caption("Track your applications from wishlist to offer.")

    apps = svc_list_applications()
    df = pd.DataFrame(apps) if apps else pd.DataFrame(
        columns=["id", "company", "role", "status", "match_score", "applied_date"]
    )

    if search and not df.empty:
        mask = (
            df["company"].str.contains(search, case=False, na=False)
            | df["role"].str.contains(search, case=False, na=False)
        )
        df = df[mask]

    render_kpi_banner(df)
    st.markdown("---")

    if df.empty:
        st.info("No applications yet. Use **➕ New Application** in the sidebar to add one.")
    else:
        render_kanban(df)


# ---------------------------------------------------------------------------
# 5.0. JOB SEARCH VIEW  (MyCareersFuture)
# ---------------------------------------------------------------------------

#def view_job_search():
#    st.title("🔎 Job Search — MyCareersFuture")
#    st.caption(
#        "Search Singapore's national job portal. Pick a result to pre-fill the "
#        "New Application form; you still control generation and saving there."
#    )

#   if not _JOB_SEARCH_AVAILABLE:
#        st.error(
#            "`job_search.py` module not found. Create it in the same directory "
#            "as `app.py` to enable this page."
#        )
#        return

#    with st.form("mcf_search_form"):
#        c1, c2 = st.columns([3, 1])
#        keywords = c1.text_input(
#            "Role / keywords",
#            placeholder="data analyst, python developer, cybersecurity...",
#        )
#        max_results = c2.number_input(
#            "Max results", min_value=10, max_value=300, value=60, step=10,
#        )

#        c3, c4 = st.columns(2)
#        sort_by = c3.selectbox(
#            "Sort by",
#            ["new_posting_date", "relevancy"],
#            format_func=lambda x: (
#                "Newest first" if x == "new_posting_date" else "Most relevant"
#            ),
#        )
#        include_desc = c4.checkbox(
#            "Include full description in preview", value=False,
#        )

#        submitted = st.form_submit_button("Search", type="primary")

#    if not submitted:
#        return
#    if not keywords.strip():
#        st.warning("Enter at least one keyword.")
#        return

#    with st.spinner(f"Searching MyCareersFuture for “{keywords}”..."):
#        try:
#            results, diag = search_mycareersfuture(
#                keywords=keywords,
#                max_results=int(max_results),
#                sort_by=sort_by,
#            )
#        except Exception as e:
#            st.error(f"Search failed: {e}")
#            return

#    for err in diag.get("errors", []):
#        st.warning(err)

#    if not results:
#        st.info("No results. Try broader keywords or a different sort order.")
#        return

#    st.success(
#        f"{len(results)} postings · {diag.get('pages_fetched', 0)} page(s) fetched"
#    )

#    for i, job in enumerate(results):
#        with st.container(border=True):
#            st.markdown(f"**{job.title}** — {job.company}")
#            meta = [job.location, job.salary or "salary not listed"]
#            if job.employment_type:
#                meta.append(job.employment_type)
#            if job.posted_date:
#                meta.append(job.posted_date)
#            st.caption(" · ".join(m for m in meta if m))

#            if job.skills:
#                st.caption("Skills: " + ", ".join(job.skills[:8]))

#            if include_desc and job.description:
#                with st.expander("Description preview"):
#                    st.write(job.description[:3000])

#            if st.button("Use this job →", key=f"mcf_use_{i}"):
#                st.session_state["app_company"] = job.company
#                st.session_state["app_role"] = job.title
#                st.session_state["app_url"] = job.url
#                st.session_state["app_desc"] = job.description
#                st.session_state["nav_to_new_app"] = True
#                st.rerun()


def view_job_search():
    st.title("🔎 Job Search — MyCareersFuture")
    st.caption("Search Singapore's job portal. Pick a job to pre-fill the New Application form.")

    col1, col2, col3 = st.columns([3, 1, 1])
    keywords = col1.text_input("Keywords", placeholder="data analyst, python developer...")
    max_results = col2.number_input("Max results", 10, 100, 40, 10)
    sort_by = col3.selectbox(
        "Sort by",
        ["new_posting_date", "relevancy"],
        format_func=lambda x: "Newest" if x == "new_posting_date" else "Relevant",
    )

    if st.button("Search", type="primary"):
        if not keywords.strip():
            st.warning("Enter some keywords.")
            return

        with st.spinner("Searching..."):
            try:
                results = search_mycareersfuture(keywords, int(max_results), sort_by=sort_by)
            except Exception as e:
                st.error(str(e))
                return

        if not results:
            st.info("No results. Try different keywords.")
            return

        st.success(f"{len(results)} results")

        for i, job in enumerate(results):
            with st.container(border=True):
                st.markdown(f"**{job.title}** — {job.company}")
                st.caption(f"{job.location} · {job.salary or 'salary not listed'} · {job.posted_date}")

                if st.button("Use this job →", key=f"use_{i}"):
                    st.session_state["app_company"] = job.company
                    st.session_state["app_role"] = job.title
                    st.session_state["app_url"] = job.url
                    st.session_state["app_desc"] = job.description
                    st.session_state["nav_to_new_app"] = True
                    st.rerun()




def view_new_application():
    st.title("➕ New Application")
    st.caption(
        "Paste the job description and upload your resume, or paste a job link "
        "below to pull details automatically."
    )

    for key in ["app_company", "app_role", "app_url", "app_desc"]:
        if key not in st.session_state:
            st.session_state[key] = ""

    with st.container(border=True):
        st.markdown("**🔗 Quick Import from Job Link**")
        col_url, col_btn = st.columns([4, 1])
        with col_url:
            quick_url = st.text_input(
                "Job Link",
                placeholder="https://...",
                label_visibility="collapsed",
                key="quick_url_input",
            )
        with col_btn:
            if st.button("Fetch Details", use_container_width=True):
                if quick_url:
                    with st.spinner("Pulling job information..."):
                        scraped = fetch_job_details_from_url(quick_url)
                        st.session_state["app_company"] = scraped["company"]
                        st.session_state["app_role"] = scraped["role"]
                        st.session_state["app_url"] = scraped["job_url"]
                        st.session_state["app_desc"] = scraped["job_description"]
                    st.success("Details pulled! Review them in the form below.")
                    st.rerun()
                else:
                    st.warning("Please paste a URL first.")

    _, _, model_label = get_active_model()
    st.info(f"Using model: **{model_label}**")

    with st.form("new_app_form", clear_on_submit=True):
        c1, c2 = st.columns(2)
        company = c1.text_input("Company Name *", value=st.session_state["app_company"])
        role = c2.text_input("Role Title *", value=st.session_state["app_role"])
        job_url = st.text_input("Job URL (optional)", value=st.session_state["app_url"])
        job_description = st.text_area(
            "Job Description", value=st.session_state["app_desc"], height=180,
        )

        resume_file = st.file_uploader("Upload Resume (PDF)", type=["pdf"])

        st.markdown("**Human-in-the-loop generation**")
        generation_choice = st.radio(
            "Choose what AI should generate",
            [
                "Resume only",
                "Cover letter only",
                "Resume and cover letter",
                "None",
            ],
            index=2,
            horizontal=True,
        )

        submitted = st.form_submit_button("🔍 Analyze & Add", type="primary")

    if submitted:
        for key in ["app_company", "app_role", "app_url", "app_desc"]:
            st.session_state[key] = ""

        company, role = clean_text(company), clean_text(role)
        if not company or not role:
            st.error("Company and Role are required.")
            return

        resume_text = extract_pdf_text(resume_file) if resume_file else ""
        match_str = ""
        tailored = ""
        cover = ""
        resume_filename = ""
        cover_letter_filename = ""
        doc_number = 0

        if resume_text and job_description:
            with st.spinner("Scoring resume against job..."):
                score = calculate_match_score(resume_text, job_description)
                match_str = f"{score:.0%}"
                st.metric("Profile Match Score", match_str)

            generate_resume = generation_choice in [
                "Resume only", "Resume and cover letter",
            ]
            generate_cover = generation_choice in [
                "Cover letter only", "Resume and cover letter",
            ]

            if generate_resume:
                with st.spinner(f"Tailoring resume with {model_label}..."):
                    tailored = tailor_resume(resume_text, job_description)

            if generate_cover:
                base_resume = tailored if tailored else resume_text
                with st.spinner(f"Drafting cover letter with {model_label}..."):
                    cover = generate_cover_letter(
                        base_resume, job_description, company, role,
                    )

            if tailored or cover:
                doc_number = get_next_doc_number()
                if tailored:
                    resume_filename = build_document_filename(
                        doc_number, role, company, "CV",
                    )
                if cover:
                    cover_letter_filename = build_document_filename(
                        doc_number, role, company, "CL",
                    )

            st.success("Document generation complete.")

            if tailored:
                with st.expander("📄 Tailored Resume (preview)", expanded=False):
                    st.text_area(
                        "Tailored Resume", tailored, height=280,
                        label_visibility="collapsed", disabled=True,
                    )

            if cover:
                with st.expander("✉️ Cover Letter (preview)", expanded=False):
                    st.text_area(
                        "Cover Letter", cover, height=280,
                        label_visibility="collapsed", disabled=True,
                    )

            app_id = svc_create_application(
                company=company,
                role=role,
                job_url=job_url,
                job_description=job_description,
                resume_text=tailored if tailored else resume_text,
                cover_letter=cover,
                match_score=match_str,
                status="Applied",
                doc_number=doc_number,
                resume_filename=resume_filename,
                cover_letter_filename=cover_letter_filename,
            )
            st.session_state["edit_app_id"] = app_id
            st.success(f"✅ Application #{app_id} saved to Tracker.")
            if tailored or cover:
                st.info(
                    "Open **✏️ Edit & Export** in the sidebar to polish the "
                    "documents and download them as PDF or DOCX."
                )
        else:
            app_id = svc_create_application(
                company=company,
                role=role,
                job_url=job_url,
                job_description=job_description,
                status="Applied",
            )
            st.info(
                f"Application #{app_id} saved without AI analysis. "
                "Upload a PDF + paste a job description to enable AI features."
            )


# ---------------------------------------------------------------------------
# 5a. EDIT & EXPORT VIEW
# ---------------------------------------------------------------------------

DOC_KINDS = {
    "Resume": {
        "text_field": "resume_text",
        "filename_field": "resume_filename",
        "doc_type": "CV",
        "prefix": "resume",
    },
    "Cover Letter": {
        "text_field": "cover_letter",
        "filename_field": "cover_letter_filename",
        "doc_type": "CL",
        "prefix": "cover",
    },
}


def render_rich_editor(initial: str, key: str, height: int = 520) -> str:
    if streamlit_lexical_extended is None:
        st.warning(
            "Rich text editor not installed "
            "(`pip install streamlit-lexical-extended`, needs Streamlit 1.51+). "
            "Using a plain text area instead."
        )
        return st.text_area(
            "Edit", initial, height=height, key=key, label_visibility="collapsed",
        )

    seeded = st.session_state.setdefault("_lexical_seeded", set())
    if key in seeded and key in st.session_state:
        seed_value = None
    else:
        seed_value = initial
        seeded.add(key)

    result = streamlit_lexical_extended(
        value=seed_value,
        placeholder="Edit your document...",
        height=height,
        debounce=500,
        key=key,
    )

    state_val = st.session_state.get(key)
    if isinstance(result, str) and result.strip():
        return result
    if isinstance(state_val, str):
        return state_val
    return initial


def _record_export(app_id: int, cfg: dict, filename: str, doc_number: int, text: str):
    try:
        svc_update_application(
            app_id,
            **{
                cfg["text_field"]: text,
                cfg["filename_field"]: filename,
                "doc_number": doc_number,
            },
        )
    except Exception as e:
        st.toast(f"Could not update the application record: {e}")


def view_edit_export():
    st.title("✏️ Edit & Export")
    st.caption("Polish the AI output, then download it as PDF or DOCX to your device.")

    apps = [a for a in svc_list_applications()
            if a.get("resume_text") or a.get("cover_letter")]
    if not apps:
        st.info("No documents yet. Create an application with a resume/cover letter first.")
        return

    apps.sort(key=lambda a: a["id"])
    labels = {a["id"]: f"#{a['id']} · {a['role']} @ {a['company']}" for a in apps}
    ids = list(labels.keys())
    pre = st.session_state.get("edit_app_id")
    app_id = st.selectbox(
        "Application", ids,
        index=ids.index(pre) if pre in ids else 0,
        format_func=lambda i: labels[i],
        key="edit_app_select",
    )
    st.session_state["edit_app_id"] = app_id
    app = svc_get_application(app_id)

    kind = st.radio("Document", list(DOC_KINDS.keys()), horizontal=True)
    cfg = DOC_KINDS[kind]
    initial = app.get(cfg["text_field"]) or ""
    if not initial.strip():
        st.info(f"No {kind.lower()} saved for this application.")
        return

    updates = {}
    doc_number = app.get("doc_number") or 0
    if not doc_number:
        doc_number = get_next_doc_number()
        updates["doc_number"] = doc_number
    filename = app.get(cfg["filename_field"]) or ""
    if not filename:
        filename = build_document_filename(
            doc_number, app["role"], app["company"], cfg["doc_type"],
        )
        updates[cfg["filename_field"]] = filename
    if updates:
        svc_update_application(app_id, **updates)

    st.caption(f"File name: `{filename}` (stored on the application record)")

    edited = render_rich_editor(initial, key=f"lexical_{app_id}_{cfg['text_field']}")

    st.markdown("---")
    if st.button("💾 Save edits", type="primary"):
        _record_export(app_id, cfg, filename, doc_number, edited)
        st.success("Edits saved.")

    try:
        pdf_bytes = markdown_to_pdf_bytes(edited)
        docx_bytes = markdown_to_docx_bytes(edited)
    except Exception as e:
        st.error(f"Export failed: {e}")
        return

    st.markdown("**⬇️ Download to your device**")
    d1, d2 = st.columns(2)
    d1.download_button(
        "Download PDF", data=pdf_bytes,
        file_name=f"{filename}.pdf", mime=MIME_PDF,
        key=f"dl_pdf_{app_id}_{cfg['prefix']}",
        on_click=_record_export,
        args=(app_id, cfg, filename, doc_number, edited),
    )
    d2.download_button(
        "Download DOCX", data=docx_bytes,
        file_name=f"{filename}.docx", mime=MIME_DOCX,
        key=f"dl_docx_{app_id}_{cfg['prefix']}",
        on_click=_record_export,
        args=(app_id, cfg, filename, doc_number, edited),
    )
    st.caption(
        "Files go to your browser's download folder. Turn on 'Ask where to save "
        "each file' in your browser settings to choose a location each time."
    )


def view_outbox():
    st.title("📥 Outbox")
    st.info("Email sync is not yet wired up. This view will show parsed job-related emails.")

    sample = pd.DataFrame({
        "Date": ["2026-10-08", "2026-10-07", "2026-10-06"],
        "Sender": ["recruiting@stripe.com", "no-reply@google.com", "hr@meta.com"],
        "Subject": ["Interview Invitation", "Application Received", "Update on your application"],
        "Detected Status": ["🟢 Interviewing", "🟡 Acknowledged", "🔴 Rejected"],
    })
    st.dataframe(sample, width="stretch", hide_index=True)

    st.markdown("**Test the classifier:**")
    body = st.text_area("Paste an email body:")
    if body:
        st.success(f"Detected status: **{classify_email_status(body)}**")


def view_analytics():
    st.title("📈 Analytics")
    apps = svc_list_applications()

    if not apps:
        st.info("No data yet. Add applications to see analytics.")
        return

    df = pd.DataFrame(apps)
    c1, c2 = st.columns(2)

    with c1:
        st.subheader("Applications by Stage")
        st.bar_chart(df["status"].value_counts())

    with c2:
        st.subheader("Applications Over Time")
        if "applied_date" in df:
            df["date"] = pd.to_datetime(df["applied_date"]).dt.date
            st.line_chart(df.groupby("date").size())


# ---------------------------------------------------------------------------
# 5b. MONITORING VIEW
# ---------------------------------------------------------------------------

def _fmt_int(n):
    return f"{int(n):,}"


def _fmt_usd(v):
    return f"${v:,.4f}" if v < 1 else f"${v:,.2f}"


def view_monitoring():
    st.title("🔍 LLM Monitoring")
    st.caption(
        "Every LLM request: tokens, cost, latency, failures. "
        "Recommendations never change anything automatically."
    )

    fc1, fc2, fc3, fc4 = st.columns([2, 2, 2, 2])
    with fc1:
        days = st.selectbox(
            "Window", [1, 7, 14, 30, 90], index=1,
            format_func=lambda d: f"Last {d} day{'s' if d > 1 else ''}",
        )
    with fc2:
        provider_filter = st.selectbox(
            "Provider",
            ["All", "groq", "ai_singapore"],
            format_func=lambda x: "AI Singapore" if x == "ai_singapore" else x,
        )
    with fc3:
        model_filter = st.selectbox("Model", ["All"] + list(MODEL_REGISTRY.keys()))
    with fc4:
        status_filter = st.selectbox("Status", ["All", "Success", "Failed"])

    rows = svc_list_llm_usage(days=days)

    if not rows:
        st.info(
            "No LLM calls recorded yet. Generate a resume or cover letter "
            "from **➕ New Application** to populate this dashboard."
        )
        _render_pricing_reference()
        return

    df = pd.DataFrame(rows)
    df["called_at"] = pd.to_datetime(df["called_at"])

    if provider_filter != "All":
        df = df[df["provider"] == provider_filter]
    if model_filter != "All":
        df = df[df["model_key"] == model_filter]
    if status_filter == "Success":
        df = df[df["success"] == 1]
    elif status_filter == "Failed":
        df = df[df["success"] == 0]

    if df.empty:
        st.warning("No data matches the current filters.")
        _render_pricing_reference()
        return

    total_requests = len(df)
    success_df = df[df["success"] == 1]
    failed_df = df[df["success"] == 0]

    total_tokens = int(df["total_tokens"].sum())
    input_tokens = int(df["input_tokens"].sum())
    cached_input_tokens = (
        int(df["cached_input_tokens"].sum())
        if "cached_input_tokens" in df.columns else 0
    )
    output_tokens = int(df["output_tokens"].sum())
    total_cost = float(df["cost_usd"].sum())

    avg_latency = (
        float(success_df["latency_seconds"].mean())
        if not success_df.empty else 0.0
    )
    p95_latency = (
        float(success_df["latency_seconds"].quantile(0.95))
        if not success_df.empty else 0.0
    )
    failure_rate = (
        len(failed_df) / total_requests * 100 if total_requests else 0.0
    )

    st.markdown("---")
    k1, k2, k3, k4 = st.columns(4)
    k1.metric("LLM Requests", _fmt_int(total_requests), f"{avg_latency:.1f}s avg latency")
    k2.metric("Total Tokens", _fmt_int(total_tokens))
    k3.metric("Uncached Input", _fmt_int(input_tokens))
    k4.metric("Cached Input", _fmt_int(cached_input_tokens))

    k5, k6, k7, k8 = st.columns(4)
    k5.metric("Output Tokens", _fmt_int(output_tokens))
    k6.metric("Estimated Cost", _fmt_usd(total_cost))
    k7.metric("Avg Latency", f"{avg_latency:.1f}s", f"P95 {p95_latency:.1f}s")
    k8.metric("Failed Requests", _fmt_int(len(failed_df)), f"{failure_rate:.1f}% of requests")

    st.markdown("---")

    c_left, c_right = st.columns(2)

    with c_left:
        st.subheader("💰 Cost over time")
        st.caption(f"Estimated USD per day · {days}-day window")
        daily = (
            df.set_index("called_at")
              .resample("D")["cost_usd"].sum()
              .reset_index()
              .rename(columns={"called_at": "date", "cost_usd": "cost"})
        )
        if daily["cost"].sum() == 0:
            st.info("All calls in this window were $0 (free-tier models).")
        else:
            st.line_chart(daily.set_index("date")["cost"])

    with c_right:
        st.subheader("🔢 Token usage over time")
        st.caption("Input (uncached + cached) vs Output per day")
        daily_tok = (
            df.set_index("called_at")
              .resample("D")[["input_tokens", "cached_input_tokens", "output_tokens"]]
              .sum()
              .reset_index()
        )
        daily_tok["Input"] = daily_tok["input_tokens"] + daily_tok["cached_input_tokens"]
        daily_tok["Output"] = daily_tok["output_tokens"]
        daily_tok = daily_tok.rename(columns={"called_at": "date"})
        st.bar_chart(daily_tok.set_index("date")[["Input", "Output"]])

    st.markdown("---")

    c_left2, c_right2 = st.columns(2)

    with c_left2:
        st.subheader("📊 Usage by model")
        by_model = (
            df.groupby("model_key")
              .agg(
                  Requests=("id", "count"),
                  UncachedInput=("input_tokens", "sum"),
                  CachedInput=("cached_input_tokens", "sum"),
                  OutputTokens=("output_tokens", "sum"),
                  TotalTokens=("total_tokens", "sum"),
                  CostUSD=("cost_usd", "sum"),
                  AvgLatency=("latency_seconds", "mean"),
              )
              .reset_index()
              .rename(columns={"model_key": "Model"})
        )
        by_model["AvgLatency"] = by_model["AvgLatency"].round(1)
        by_model["CostUSD"] = by_model["CostUSD"].round(4)
        st.dataframe(by_model, width="stretch", hide_index=True)

    with c_right2:
        st.subheader("💵 Cost by model")
        cost_by_model = by_model.set_index("Model")["CostUSD"]
        if cost_by_model.sum() == 0:
            st.info("All selected models are free-tier ($0).")
        else:
            st.bar_chart(cost_by_model)

    st.markdown("---")

    st.subheader("⚙️ Usage by operation")
    by_op = (
        df.groupby("operation")
          .agg(
              Requests=("id", "count"),
              TotalTokens=("total_tokens", "sum"),
              CostUSD=("cost_usd", "sum"),
              AvgLatency=("latency_seconds", "mean"),
          )
          .reset_index()
          .rename(columns={"operation": "Operation"})
    )
    by_op["AvgLatency"] = by_op["AvgLatency"].round(1)
    by_op["CostUSD"] = by_op["CostUSD"].round(4)
    st.dataframe(by_op, width="stretch", hide_index=True)

    st.markdown("---")

    if not failed_df.empty:
        st.subheader("❌ Recent failures")
        fail_view = (
            failed_df[["called_at", "model_key", "operation",
                       "latency_seconds", "error_message"]]
            .head(10)
            .rename(columns={
                "called_at": "Time",
                "model_key": "Model",
                "operation": "Op",
                "latency_seconds": "Latency (s)",
                "error_message": "Error",
            })
        )
        st.dataframe(fail_view, width="stretch", hide_index=True)

    st.markdown("---")

    with st.expander("📜 Raw call log (last 100)"):
        raw = df.head(100)[[
            "called_at", "provider", "model_key", "operation",
            "input_tokens", "cached_input_tokens", "output_tokens",
            "total_tokens", "latency_seconds", "cost_usd", "success",
        ]]
        st.dataframe(raw, width="stretch", hide_index=True)

    with st.expander("🗑 Danger zone"):
        if st.button("Clear all LLM usage logs", type="secondary"):
            with db_session() as db:
                db.query(LLMUsage).delete()
                db.commit()
            st.success("LLM usage logs cleared.")
            st.rerun()

    st.markdown("---")
    _render_pricing_reference()


def _render_pricing_reference():
    st.subheader("💲 Model pricing reference")
    st.caption(
        "Per 1M tokens (USD). Groq prices per user's plan; "
        "AI Singapore Sea-Lion is free-tier."
    )

    pricing_rows = []
    for key, entry in MODEL_REGISTRY.items():
        p = get_pricing(entry["model_id"])
        pricing_rows.append({
            "Model Key": key,
            "Provider": entry["provider"],
            "Model ID": entry["model_id"],
            "Input $/1M": p["input"],
            "Cached Input $/1M": p["cached_input"],
            "Output $/1M": p["output"],
            "Context": p["context_window"],
            "Max Output": p["max_output"],
        })
    st.dataframe(pd.DataFrame(pricing_rows), width="stretch", hide_index=True)

    st.info(
        "**Cost formula:**  \n"
        "`Cost = (uncached_input_tokens / 1,000,000) × input_price "
        "+ (cached_input_tokens / 1,000,000) × cached_input_price "
        "+ (output_tokens / 1,000,000) × output_price`"
    )


# ---------------------------------------------------------------------------
# 5c. SETTINGS VIEW
# ---------------------------------------------------------------------------

def view_settings():
    st.title("⚙️ Settings")

    provider, model_id, label = get_active_model()

    st.markdown("**Active Model**")
    st.code(
        f"Display key : {label}\n"
        f"Provider    : {provider}\n"
        f"Model ID    : {model_id}",
        language="bash",
    )

    st.markdown("**API Keys Detected**")
    st.write(f"- GROQ_API_KEY: {'✅ set' if GROQ_API_KEY else '❌ missing'}")
    st.write(f"- SEALION_API_KEY: {'✅ set' if SEALION_API_KEY else '❌ missing'}")
    st.write(f"- HF_TOKEN: {'✅ set' if HF_TOKEN else '⚪ not set (optional)'}")
    st.write(f"- SUPABASE_DB_URL: {'✅ set' if SUPABASE_DB_URL else '❌ missing'}")

    st.markdown("**UI Components**")
    st.write(
        f"- streamlit-lexical-extended (rich editor): "
        f"{'✅ installed' if streamlit_lexical_extended else '❌ missing'}"
    )
    st.write(
        f"- job_search (MyCareersFuture adapter): "
        f"{'✅ installed' if _JOB_SEARCH_AVAILABLE else '❌ missing'}"
    )

    st.markdown("**Test Direct Calls**")
    if provider == "groq":
        if st.button("Test Groq direct call"):
            with st.spinner(f"Calling Groq ({model_id})..."):
                st.info(call_groq_direct("Reply with: 'Groq OK'"))
    elif provider == "ai_singapore":
        if st.button("Test AI Singapore direct call"):
            with st.spinner(f"Calling AI Singapore ({model_id})..."):
                st.info(call_ai_singapore_direct("Reply with: 'AI Singapore OK'"))

    st.markdown("---")
    st.markdown("**Available Models**")
    for key, entry in MODEL_REGISTRY.items():
        marker = " ← active" if key == label else ""
        st.markdown(
            f"- **{key}**{marker}  \n"
            f"  `{entry['model_id']}`  \n"
            f"  {entry['description']}"
        )

    st.markdown("---")
    st.markdown("**Danger Zone**")
    col1, col2 = st.columns(2)
    with col1:
        if st.button("🗑 Clear all applications", type="secondary"):
            with db_session() as db:
                db.query(Application).delete()
                db.commit()
            st.success("All applications cleared.")
            st.rerun()
    with col2:
        if st.button("🗑 Clear LLM usage logs", type="secondary"):
            with db_session() as db:
                db.query(LLMUsage).delete()
                db.commit()
            st.success("LLM usage logs cleared.")
            st.rerun()


# ---------------------------------------------------------------------------
# 6. MAIN
# ---------------------------------------------------------------------------

def main():
    if "selected_model_key" not in st.session_state:
        st.session_state["selected_model_key"] = DEFAULT_MODEL_KEY

    # Honor cross-view navigation requests (e.g. "Use this job" → New Application)
    if st.session_state.pop("nav_to_new_app", False):
        st.session_state["nav_radio"] = "➕ New Application"

    nav, search = render_sidebar()

    if nav == "📊 Tracker":
        view_tracker(search)
    elif nav == "🔎 Job Search":
        view_job_search()
    elif nav == "➕ New Application":
        view_new_application()
    elif nav == "✏️ Edit & Export":
        view_edit_export()
    elif nav == "📥 Outbox":
        view_outbox()
    elif nav == "📈 Analytics":
        view_analytics()
    elif nav == "🔍 Monitoring":
        view_monitoring()
    elif nav == "⚙️ Settings":
        view_settings()


if __name__ == "__main__":
    main()
